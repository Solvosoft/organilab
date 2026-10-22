"""Estilo común de las hojas de cálculo de los reportes (XLSX, XLS y ODS)."""

import zipfile

import openpyxl
import pyexcel
from django.test import SimpleTestCase

from laboratory.report_utils import (
    BRAND,
    ExcelGraphBuilder,
    infer_spreadsheet_roles,
    sheet_name,
)

FORMATOS = ("xlsx", "xls", "ods")


def roles(data):
    return [role for role, _cells in infer_spreadsheet_roles(data)]


def leer(contenido, formato):
    """Filas de la primera hoja, sin las celdas vacías del final."""
    hoja = pyexcel.get_book(file_type=formato, file_content=contenido.getvalue())
    filas = []
    for fila in hoja[0].to_array():
        fila = ["" if c is None else c for c in fila]
        while fila and fila[-1] == "":
            fila.pop()
        filas.append(fila)
    return filas


class RolesTest(SimpleTestCase):
    def test_tabla_simple(self):
        data = [["reporte"], ["Nombre", "CAS"], ["Cal", "1305-78-8"], ["Sal", None]]
        self.assertEqual(roles(data), ["title", "header", "cell", "cell"])

    def test_secciones_por_laboratorio(self):
        data = [
            ["reporte"],
            ["Lab A"], ["Estante", "Objeto"], ["E1", "Cal"],
            ["Lab B"], ["Estante", "Objeto"],
        ]
        self.assertEqual(
            roles(data), ["title", "section", "header", "cell", "section", "header"]
        )

    def test_dato_valor_antes_de_la_tabla(self):
        data = [
            ["reporte"], [""],
            ["Laboratorio", "Lab A"], ["Teléfono", "2222"],
            ["Nombre", "CAS", "Cantidad"], ["Cal", "1305-78-8", 3],
        ]
        self.assertEqual(
            roles(data), ["title", "blank", "label", "label", "header", "cell"]
        )

    def test_bloques_de_movimientos(self):
        data = [
            ["Lab | Cal"], ["Diferencia: 2 Litros"],
            ["Usuario", "Día", "Diferencia"], ["Ana", "01/02", 2], [], [],
            ["Lab | Sal"], ["Diferencia: 0 Litros"], [], [],
        ]
        self.assertEqual(
            roles(data),
            ["title", "section", "header", "cell", "blank", "blank",
             "section", "section", "blank", "blank"],
        )

    def test_fila_que_no_es_lista(self):
        self.assertEqual(roles([["reporte"], "texto suelto"]), ["title", "section"])

    def test_nombre_de_hoja(self):
        self.assertEqual(sheet_name("informe: [lab]/2026?"), "informe   lab  2026")
        self.assertEqual(len(sheet_name("x" * 40)), 31)
        self.assertEqual(sheet_name(""), "Reporte")


class SaveOdsTest(SimpleTestCase):
    DATA = [
        ["reporte-prueba"],
        ["Nombre", "Cantidad", "Nota"],
        ["Cal", 1.5, None],
        ["Sal", 2, "ácido"],
    ]

    def test_el_contenido_no_cambia_en_ningun_formato(self):
        esperado = [["reporte-prueba"], ["Nombre", "Cantidad", "Nota"], ["Cal", 1.5], ["Sal", 2, "ácido"]]
        for formato in FORMATOS:
            with self.subTest(formato=formato):
                salida = ExcelGraphBuilder().save_ods([list(f) for f in self.DATA], format_type=formato)
                self.assertEqual(leer(salida, formato), esperado)

    def test_celda_vacia_en_ods_no_falla(self):
        # pyexcel-ods3 no aceptaba None: el ODS de «Objetos» moría con KeyError.
        salida = ExcelGraphBuilder().save_ods([["r"], ["A", "B"], [None, 1]], format_type="ods")
        self.assertEqual(leer(salida, "ods")[2], ["", 1])

    def test_fila_nula_y_texto_suelto_en_todos_los_formatos(self):
        data = [["reporte"], None, "texto suelto", ["A", None, 1]]
        for formato in FORMATOS:
            with self.subTest(formato=formato):
                salida = ExcelGraphBuilder().save_ods(data, format_type=formato)
                self.assertEqual(leer(salida, formato), [["reporte"], [], ["texto suelto"], ["A", "", 1]])

    def test_estilos_del_xlsx(self):
        salida = ExcelGraphBuilder().save_ods([list(f) for f in self.DATA], format_type="xlsx")
        ws = openpyxl.load_workbook(salida).worksheets[0]
        self.assertEqual(ws.title, "reporte-prueba")
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws["A2"].fill.fgColor.rgb[-6:], BRAND)
        self.assertTrue(ws["A3"].alignment.wrap_text)
        self.assertEqual(ws.freeze_panes, "A3")
        self.assertEqual(ws.auto_filter.ref, "A2:C4")

    def test_el_titulo_no_sale_cortado(self):
        data = [["reporte-prueba"], ["Lab A"], ["X", "Y"], [1, 2]]
        ws = openpyxl.load_workbook(
            ExcelGraphBuilder().save_ods(data, format_type="xlsx")
        ).worksheets[0]
        self.assertEqual(ws.row_dimensions[1].height, 24)
        self.assertEqual(ws.row_dimensions[2].height, 18)
        ods = ExcelGraphBuilder().save_ods(data, format_type="ods")
        with zipfile.ZipFile(ods) as archivo:
            content = archivo.read("content.xml").decode()
        self.assertIn('style:row-height="24pt"', content)
        self.assertIn('table:style-name="report_row_title"', content)
        xls = ExcelGraphBuilder().save_ods(data, format_type="xls")
        self.assertEqual(leer(xls, "xls")[0], ["reporte-prueba"])

    def test_varias_tablas_no_fija_encabezado(self):
        data = [["r"], ["Lab A"], ["X", "Y"], [1, 2], ["Lab B"], ["X", "Y"], [3, 4]]
        ws = openpyxl.load_workbook(
            ExcelGraphBuilder().save_ods(data, format_type="xlsx")
        ).worksheets[0]
        self.assertIsNone(ws.freeze_panes)

    def test_otros_formatos_siguen_igual(self):
        salida = ExcelGraphBuilder().save_ods([["a", "b"], [1, 2]], format_type="csv")
        self.assertIn("a,b", salida.getvalue())
