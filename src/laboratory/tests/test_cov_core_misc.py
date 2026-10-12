from io import BytesIO
from types import SimpleNamespace
from unittest import mock

from django.contrib.auth.models import Group, User
from django.core.exceptions import ValidationError
from django.db import models as dj_models
from django.test import RequestFactory, SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from djreservation.models import Product, Reservation
from openpyxl import load_workbook

from laboratory import dataconfig, lab_or_org_request_notifications as notificaciones
from laboratory import shelf_utils, utils_upload_reatives as carga, validators
from laboratory.catalog import GTForeignKey
from laboratory.gtcharts import BaseChart
from laboratory.logsustances import log_object_add_change
from laboratory.models import (
    Catalog,
    LabOrOrgRequest,
    Laboratory,
    MaterialCapacity,
    Object,
    ObjectLogChange,
    ObjectMaximumLimit,
    OrganizationStructure,
    ReactiveLimit,
    Shelf,
    ShelfObject,
    ShelfObjectObservation,
)
from laboratory.report_utils import ExcelGraphBuilder
from laboratory.shelfobject_container.serializers import ContainerFilter, ContainerListSerializer
from sga.models import SubstanceCharacteristics


class DataconfigToleranteTest(SimpleTestCase):

    def test_celdas_de_formatos_historicos(self):
        self.assertEqual(dataconfig._cell_to_pks(None), [])
        self.assertEqual(dataconfig._cell_to_pks(True), [])
        self.assertEqual(dataconfig._cell_to_pks(3.0), [3])
        self.assertEqual(dataconfig._cell_to_pks("1,,x, 4"), [1, 4])
        self.assertEqual(dataconfig._cell_to_pks({"pk": 1}), [])

    def test_texto_no_deserializable(self):
        self.assertIs(dataconfig._deserialize(5), dataconfig.UNPARSEABLE)
        self.assertEqual(dataconfig.parse(5), [])

    def test_parse_strict(self):
        self.assertEqual(dataconfig.parse_strict('[[["1,,2"]]]'), [[[1, 2]]])
        with self.assertRaises(dataconfig.DataconfigInvalid):
            dataconfig.parse_strict("[[[1.5]]]")

    def test_matriz_del_servicio(self):
        servicio = dataconfig.DataconfigService(SimpleNamespace(dataconfig="[[[1],[2,2]]]"))
        self.assertEqual(servicio.get_matrix(), [[[1], [2]]])


class ValidadoresYCatalogoTest(SimpleTestCase):

    def test_formula_molecular(self):
        self.assertTrue(validators.isValidate_molecular_formula("H2O"))
        self.assertFalse(validators.isValidate_molecular_formula(""))
        validators.validate_molecular_formula("H2O")
        with self.assertRaises(ValidationError):
            validators.validate_molecular_formula("")

    def test_foreignkey_sin_modelo_resuelto(self):
        campo = GTForeignKey("laboratory.Catalog", on_delete=dj_models.CASCADE)
        with self.assertRaises(ValueError):
            campo.formfield()


class ExcelGraphBuilderTest(SimpleTestCase):

    def test_tabla_estilos_y_guardado(self):
        builder = ExcelGraphBuilder()
        builder.add_table([["Nombre", "Cantidad"], ["Acetona", 3]], "Inventario")
        builder.style_header()
        builder.autosize_columns(max_width=5)
        builder.format_border_cell(3, 2)
        salida = builder.save()
        self.assertIsInstance(salida, BytesIO)
        hoja = load_workbook(BytesIO(salida.getvalue())).active
        self.assertEqual(hoja["A1"].value, "Inventario")
        self.assertEqual(hoja["A3"].value, "Acetona")
        self.assertTrue(hoja["A1"].font.bold)
        self.assertEqual(hoja.column_dimensions["A"].width, 5)
        self.assertEqual(hoja["B2"].border.top.style, "thin")
        self.assertEqual(builder.row_max, 3)
        otra = BytesIO()
        self.assertIs(builder.save(otra), otra)

    def test_booleanos_legibles(self):
        builder = ExcelGraphBuilder()
        self.assertEqual([builder.safe_bool(v) for v in (True, "1", "true")], ["Yes"] * 3)
        self.assertEqual(builder.safe_bool(None), "No")


class OrdenDeEstantesTest(SimpleTestCase):

    def test_preserva_el_orden_pedido(self):
        estantes = [SimpleNamespace(pk=1), SimpleNamespace(pk=2), SimpleNamespace(pk=3)]
        self.assertEqual([s.pk for s in shelf_utils.preserve_order(["3", "x", 1, 9], estantes)], [3, 1])


class NotificacionesDeSolicitudesTest(TestCase):

    def setUp(self):
        self.solicitante = User.objects.create_user("solicita", email="solicita@example.org", first_name="Ana")
        revisor = User.objects.create_user("revisa", email="revisa@example.org")
        sin_correo = User.objects.create_user("sincorreo", email="")
        grupo, _ = Group.objects.get_or_create(name="RegisterOrganization")
        grupo.user_set.add(revisor, sin_correo)
        self.solicitud = LabOrOrgRequest(name="Lab nuevo", entity_type="lab", requested_by=self.solicitante)

    def test_aviso_de_solicitud_nueva_a_revisores(self):
        """Un revisor sin correo en el grupo no debe dejar sin aviso a los demás: antes
        `exclude(user__email="")` sobre el grupo lo descartaba entero."""
        with mock.patch.object(notificaciones, "send_email_from_template") as enviar:
            notificaciones.notify_request_created(self.solicitud)
        enviar.assert_called_once()
        codigo, correo = enviar.call_args.args
        self.assertEqual((codigo, correo), ("lab_or_org_request_created", "revisa@example.org"))
        self.assertIn("Lab nuevo", enviar.call_args.kwargs["context"]["subject"])
        self.assertIn("Ana", enviar.call_args.kwargs["context"]["body"])

    def test_aviso_de_cambio_de_estado(self):
        self.solicitud.status = LabOrOrgRequest.STATUS_APPROVED
        with mock.patch.object(notificaciones, "send_email_from_template") as enviar:
            notificaciones.notify_request_status_changed(self.solicitud)
            self.solicitud.requested_by.email = ""
            notificaciones.notify_request_status_changed(self.solicitud)
        enviar.assert_called_once()
        self.assertEqual(enviar.call_args.args[1], "solicita@example.org")

    def test_un_fallo_de_correo_no_interrumpe(self):
        with mock.patch.object(notificaciones, "send_email_from_template", side_effect=RuntimeError("smtp")), \
                mock.patch.object(notificaciones.logger, "error") as registrar:
            notificaciones.notify_request_created(self.solicitud)
        registrar.assert_called_once()


class BaseInventarioTest(TestCase):
    fixtures = ["laboratory_data.json"]

    def setUp(self):
        self.user = User.objects.get(username="admin")
        self.org = OrganizationStructure.objects.get(pk=1)
        self.lab = Laboratory.objects.get(pk=1)


class CargaDeReactivosTest(BaseInventarioTest):

    def test_limites_y_unidades(self):
        objeto = Object.objects.get(pk=1)
        carga.create_reactive_limits(self.lab, objeto, 10, "mililitros")
        limite = ReactiveLimit.objects.get(laboratory=self.lab, object=objeto)
        self.assertEqual((limite.maximum_limit, limite.minimum_limit, limite.measurement_unit_id), (10, 0, 63))

    def test_conversion_segun_el_estante(self):
        con_unidad = Shelf.objects.get(pk=1)  # mililitros
        self.assertEqual(carga.validate_shelf(con_unidad, "litros", 2), 2000)
        self.assertIsNone(carga.validate_shelf(Shelf.objects.get(pk=2), "litros", 2))

    def test_reactivo_por_cas_se_crea_una_vez(self):
        nuevo = carga.get_reactive_by_cas_or_name(self.user, "64-17-5", "etanol", "C2H6O", self.org)
        self.assertEqual((nuevo.name, nuevo.type, nuevo.is_pure), ("Etanol", "0", True))
        self.assertTrue(SubstanceCharacteristics.objects.filter(object_related=nuevo, cas_id_number="64-17-5").exists())
        otra_vez = carga.get_reactive_by_cas_or_name(self.user, "64-17-5", "ETANOL", "C2H6O", self.org)
        self.assertEqual(otra_vez, nuevo)

    def test_material_contenedor_se_crea_una_vez(self):
        envase = carga.get_or_create_material(self.user, "frasco ámbar", 500, "mililitros", self.org)
        capacidad = MaterialCapacity.objects.get(object=envase)
        self.assertEqual((capacidad.capacity, capacidad.capacity_measurement_unit_id), (500, 63))
        self.assertTrue(envase.is_container)
        self.assertEqual(carga.get_or_create_material(self.user, "Frasco ámbar", 500, "mililitros", self.org), envase)

    @override_settings(ALLOWED_HOSTS=["testserver"])
    def test_crear_shelfobject_registra_observacion_bitacora_y_qr(self):
        request = RequestFactory().get("/")
        request.user = self.user
        datos = {
            "shelf": Shelf.objects.get(pk=1),
            "object": Object.objects.get(pk=4),
            "quantity": 1,
            "measurement_unit": Catalog.objects.get(pk=64),
            "in_where_laboratory": self.lab,
        }
        so = carga.create_shelfobject(datos, self.org.pk, request)
        self.assertTrue(ShelfObjectObservation.objects.filter(shelf_object=so).exists())
        self.assertTrue(ObjectLogChange.objects.filter(object=so.object, laboratory=self.lab, new_value=1).exists())
        so.refresh_from_db()
        self.assertTrue(so.shelf_object_url)


class BitacoraYSerializadoresTest(BaseInventarioTest):

    def test_bitacora_con_organizacion_como_pk(self):
        so = ShelfObject.objects.get(pk=1)
        log_object_add_change(self.user, self.lab.pk, so, 0, 5, "Ingreso", None, "F-9", create=True, organization=self.org.pk)
        log = ObjectLogChange.objects.get(bill="F-9")
        self.assertEqual((log.organization_where_action_taken_id, log.diff_value), (self.org.pk, 5))

    def test_m2m_de_catalogo_con_inicial_invocable(self):
        campo = SubstanceCharacteristics._meta.get_field("white_organ")
        formfield = campo.formfield(initial=lambda: Catalog.objects.filter(pk__in=[114, 115]))
        self.assertEqual(sorted(formfield.initial), [114, 115])

    def test_filtro_y_creador_de_contenedores(self):
        sin_nombre = User.objects.create_user("solo_usuario")
        ShelfObject.objects.filter(pk=1).update(created_by=sin_nombre)
        filtro = ContainerFilter()
        encontrados = filtro.filter_user(ShelfObject.objects.all(), "created_by", "solo_us")
        self.assertEqual(list(encontrados), [ShelfObject.objects.get(pk=1)])
        serializer = ContainerListSerializer()
        self.assertEqual(serializer.get_created_by(ShelfObject.objects.get(pk=1)), "solo_usuario")


class SenalesDeInventarioTest(BaseInventarioTest):

    @override_settings(DEBUG=True)
    def test_aviso_de_limite_usa_http_en_debug(self):
        so = ShelfObject.objects.get(pk=1)
        so.limit_quantity = so.quantity + 10
        with mock.patch("laboratory.signals.send_email_from_template") as enviar:
            so.save()
        self.assertTrue(enviar.called)
        self.assertTrue(enviar.call_args.kwargs["context"]["blockurl"].startswith("http://"))


class VistasDeAlcanceWebTest(BaseInventarioTest):

    def setUp(self):
        super().setUp()
        self.superusuario = User.objects.get(username="tuwp")
        self.superusuario.is_superuser = True
        self.superusuario.save()
        self.org.users.add(self.superusuario)
        self.client.force_login(self.superusuario)

    def test_contenedores_sin_permiso_en_el_laboratorio(self):
        """Pertenecer a la organización no basta: sin rol en ella ni en el laboratorio es 404."""
        url = reverse("laboratory:shelf_containers", kwargs={"org_pk": self.org.pk, "lab_pk": self.lab.pk})
        # HandleErrorMiddleware convierte el 404 de una página en redirección a la vista de error.
        self.assertRedirects(self.client.get(url, {"shelf": 1}), reverse("error_view") + "?status=404", fetch_redirect_response=False)
        api = reverse("laboratory:api-container-in-shelf-list", kwargs={"org_pk": self.org.pk, "lab_pk": self.lab.pk, "shelf": 1})
        self.assertEqual(self.client.get(api, HTTP_ACCEPT="application/json").status_code, 404)

    def test_protocolo_invalido_vuelve_al_formulario(self):
        admin = self.user
        admin.is_superuser = True
        admin.save()
        self.client.force_login(admin)
        url = reverse("laboratory:protocol_create", kwargs={"org_pk": self.org.pk, "lab_pk": self.lab.pk})
        respuesta = self.client.post(url, {"file": ""})
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.context["form"].errors)

    def test_reserva_de_shelfobject(self):
        url = reverse("laboratory:object_reservation", kwargs={"modelpk": 1})
        sin_reserva = self.client.get(url)
        self.assertRedirects(sin_reserva, reverse("add_user_reservation"), fetch_redirect_response=False)

        reserva = Reservation.objects.create(user=self.superusuario, reserved_end_date="2030-01-01T00:00Z")
        self.client.cookies["reservation"] = str(reserva.pk)
        self.assertEqual(self.client.get(url).status_code, 200)
        respuesta = self.client.post(url, {"amount": 2, "model_instance": 1, "available_amount": 23})
        self.assertEqual(respuesta.status_code, 302)
        producto = Product.objects.get(reservation=reserva)
        self.assertEqual((producto.amount, producto.object_id), (2, 1))

    def test_grafico_de_maximos_mensuales(self):
        ObjectMaximumLimit.objects.create(laboratory=self.lab, object_id=1, quantity=4, measurement_unit_id=63)
        anio = ObjectMaximumLimit.objects.first().created_at.year
        url = reverse("objectlimitschart-detail", kwargs={"pk": self.org.pk})
        respuesta = self.client.get(url, {"laboratory": self.lab.pk, "object": 1, "years": anio})
        self.assertEqual(respuesta.status_code, 200)
        datos = respuesta.json()
        self.assertEqual(len(datos["data"]["labels"]), 12)
        self.assertIn(4, [float(v) for v in datos["data"]["datasets"][0]["data"]])

        sin_filtros = self.client.get(url, {"laboratory": self.lab.pk})
        self.assertEqual(sin_filtros.status_code, 200)
        self.assertEqual(self.client.get(reverse("objectlimitschart-list")).status_code, 404)

    def test_filtros_base_del_grafico_con_datos_invalidos(self):
        grafico = BaseChart()
        grafico.request = SimpleNamespace(query_params={})
        grafico.get_extra_filters()
        self.assertEqual(grafico.filters, {})
        self.assertFalse(grafico.no_labs)
        grafico.index = len(grafico.colors) - 1
        self.assertEqual(grafico.get_color(), "rgb(%s)" % grafico.colors[0])
