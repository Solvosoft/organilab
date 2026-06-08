import json

from django.core.management.base import BaseCommand
from django.db.models import Count
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from laboratory.models import Object, ShelfObject


class Command(BaseCommand):
    help = "Find Object records with duplicate names and export to Excel"

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            type=str,
            default="duplicate_objects.xlsx",
            help="Output Excel file path (default: duplicate_objects.xlsx)",
        )
        parser.add_argument(
            "--json-output",
            type=str,
            default="duplicates_without_shelfobject.json",
            help="Output JSON file for duplicates without shelfobject",
        )

    def handle(self, *args, **options):
        qs = Object.objects.filter(type=Object.REACTIVE)

        duplicates = (
            qs.values("name")
            .annotate(count=Count("id"))
            .filter(count__gte=2)
            .order_by("-count")
        )

        if not duplicates.exists():
            self.stdout.write(self.style.SUCCESS("No duplicate names found."))
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Duplicados"

        headers = [
            "Nombre",
            "Cantidad",
            "ID",
            "Tipo",
            "Código",
            "Organización",
            "Tiene Security Sheet",
            "Tiene shelfobject",
        ]
        header_font = Font(bold=True)
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

        row = 2
        total_duplicates = 0

        for dup in duplicates:
            objects = qs.filter(name=dup["name"]).select_related(
                "organization", "sustancecharacteristics"
            )
            first_row = True

            for obj in objects:
                type_display = str(dict(Object.TYPE_CHOICES).get(obj.type, obj.type))
                org_name = obj.organization.name if obj.organization else "N/A"

                has_security_sheet = "No"
                if (
                    hasattr(obj, "sustancecharacteristics")
                    and obj.sustancecharacteristics
                ):
                    if obj.sustancecharacteristics.security_sheet:
                        has_security_sheet = "Sí"

                has_shelfobject = "No"
                if ShelfObject.objects.filter(object=obj).exists():
                    has_shelfobject = "Sí"
                ws.cell(row=row, column=1, value=dup["name"] if first_row else "")
                ws.cell(row=row, column=2, value=dup["count"] if first_row else "")
                ws.cell(row=row, column=3, value=obj.pk)
                ws.cell(row=row, column=4, value=type_display)
                ws.cell(row=row, column=5, value=obj.code or "N/A")
                ws.cell(row=row, column=6, value=org_name)
                ws.cell(row=row, column=7, value=has_security_sheet)
                ws.cell(row=row, column=8, value=has_shelfobject)

                for c in range(1, 9):
                    ws.cell(row=row, column=c).border = border

                row += 1
                first_row = False

            total_duplicates += dup["count"]

        for col in range(1, 9):
            ws.column_dimensions[get_column_letter(col)].width = 25

        last_row = row - 1
        ws.auto_filter.ref = f"A1:H{last_row}"

        output_path = options["output"]
        wb.save(output_path)

        duplicates_without_shelfobject = {}
        for dup in duplicates:
            objects = qs.filter(name=dup["name"])
            objects_without_shelf = []
            for obj in objects:
                if not ShelfObject.objects.filter(object=obj).exists():
                    objects_without_shelf.append(obj.pk)
            if objects_without_shelf:
                duplicates_without_shelfobject[dup["name"]] = objects_without_shelf

        json_output_path = options["json_output"]
        with open(json_output_path, "w", encoding="utf-8") as f:
            json.dump(duplicates_without_shelfobject, f, ensure_ascii=False, indent=2)

        self.stdout.write(
            self.style.SUCCESS(
                f"Exported {len(duplicates_without_shelfobject)} names without shelfobject to {json_output_path}"
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                f"Exported {len(duplicates)} duplicate names ({total_duplicates} records) to {output_path}"
            )
        )
