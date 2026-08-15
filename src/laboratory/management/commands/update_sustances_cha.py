import openpyxl
from django.core.management import CommandError
from django.core.management.base import BaseCommand

from laboratory.models import Object
from sga.models import (
    DangerIndication,
    WarningClass,
    WarningWord,
    PrudenceAdvice,
    SubstanceCharacteristics,
)
import re


class Command(BaseCommand):

    help = "Actualiza características de sustancias desde un archivo Excel"

    def add_arguments(self, parser):
        parser.add_argument(
            "fds",
            type=str,
            help="Ruta al archivo Excel (.xlsx) con los datos a procesar",
        )

    def initial_data(self):
        self.h_codes = {
            "H360F": {
                "wc": WarningClass.objects.filter(
                    name="Toxicidad para la reproducción"
                ),
                "ww": "Peligro",
                "prudence_advice": PrudenceAdvice.objects.filter(
                    code__in=["P203", "P280", "P405"]
                ),
                "description": "Puede perjudicar a la fertilidad",
            },
            "H360D": {
                "wc": WarningClass.objects.filter(
                    name="Toxicidad para la reproducción"
                ),
                "ww": "Peligro",
                "prudence_advice": PrudenceAdvice.objects.filter(
                    code__in=["P202", "P280", "P308 + P313"]
                ),
                "description": "Puede dañar al feto",
            },
            "H360FD": {
                "wc": WarningClass.objects.filter(
                    name="Toxicidad para la reproducción"
                ),
                "ww": "Peligro",
                "prudence_advice": PrudenceAdvice.objects.filter(
                    code__in=["P201", "P280", "P308 + P313"]
                ),
                "description": "Puede perjudicar a la fertilidad. Puede dañar al feto",
            },
        }

    def create_danger_indication(self):
        for h_code, data in self.h_codes.items():
            ww = WarningWord.objects.filter(name=data["ww"]).first()
            obj, created = DangerIndication.objects.get_or_create(
                code=h_code,
                description=data["description"],
                warning_words=ww,
                danger_type="health",
            )
            obj.warning_category.add(*data["wc"])
            obj.warning_class.add(*data["wc"])
            obj.prudence_advice.add(*data["prudence_advice"])
            obj.save()

    def parsear_instruccion(self, texto):
        """
        Analiza una instrucción con formato:
        'Agregar HXXX, HYYY ; Quitar HAAA, HBBB, ...'
        y retorna un diccionario con las listas separadas.
        """
        resultado = {"Agregar": [], "Quitar": []}

        # Separar en dos partes: "Agregar ..." y "Quitar ..."
        partes = re.split(r";", texto, maxsplit=1)

        for parte in partes:
            parte = parte.strip()

            # Detectar si es sección AGREGAR
            if re.match(r"(?i)Agregar", parte):
                contenido = re.sub(r"(?i)^Agregar\s*", "", parte).strip()
                codigos = re.findall(r"H\d+[A-Za-z]*", contenido)
                resultado["Agregar"] = codigos

            # Detectar si es sección QUITAR
            if re.match(r"(?i)Quitar", parte):
                contenido = re.sub(r"(?i)^Quitar\s*", "", parte).strip()
                codigos = re.findall(r"H\d+[A-Za-z]*", contenido)
                resultado["Quitar"] = codigos

        return resultado

    def read_docs(self, documento):
        try:
            wb = openpyxl.load_workbook(documento)
        except FileNotFoundError:
            raise CommandError(f"Archivo no encontrado: {documento}")
        except Exception as e:
            raise CommandError(f"Error al leer el archivo: {e}")
        ws = wb.active
        i = 0
        for fila in ws.iter_rows(min_row=2, values_only=True):
            objs = Object.objects.filter(name=fila[0])
            h_codes = fila[2].split(",")
            resultado = self.parsear_instruccion(fila[4])
            h_codes = [h_code.strip() for h_code in h_codes]
            for obj in objs:
                susta = SubstanceCharacteristics.objects.filter(
                    object_related=obj
                ).first()
                if susta:

                    i += 1
                    if (
                        len(resultado["Agregar"]) > 0
                        or len(resultado["Quitar"]) > 0
                        or fila[3] != "Cambiar por 'Ninguno'"
                    ):
                        susta.h_code.clear()
                        susta.h_code.add(*h_codes)
                        susta.save()
                else:
                    print(f"No encontro objeto {fila[0]}")

    def update_sustances_hcoded(self, documento):
        # documento = "laboratory/management/commands/ajustes_a_hojas.xlsx"
        try:
            wb = openpyxl.load_workbook(documento)
        except FileNotFoundError:
            raise CommandError(f"Archivo no encontrado: {documento}")
        except Exception as e:
            raise CommandError(f"Error al leer el archivo: {e}")
        ws = wb.active
        i = 0
        for fila in ws.iter_rows(min_row=2, values_only=True):
            obj = Object.objects.filter(name=fila[0])
            if obj.exists():
                for obj_susta in obj:
                    susta = SubstanceCharacteristics.objects.filter(
                        object_related=obj_susta
                    ).first()
                    if susta is None:
                        print(f"No encontro caracteristicas de {fila[0]}")
                        continue
                    if fila[1] != None:
                        hcode = fila[1].split(",")
                        susta.h_code.add(*hcode)

                    if fila[2] != None:
                        hcode = fila[2].split(",")
                        susta.h_code.remove(*hcode)

                    susta.save()

                i += 1

    def handle(self, *args, **options):
        self.initial_data()
        # self.create_danger_indication()
        self.read_docs(options["fds"])
