import uuid

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.cache import cache
from django.forms import formset_factory
from django.shortcuts import get_object_or_404, render, redirect
from rest_framework import status
from rest_framework.response import Response

from auth_and_perms.organization_utils import (
    user_is_allowed_on_organization,
    organization_can_change_laboratory,
)
from laboratory.api.serializers import LoadArchiveSerializer
from laboratory.forms import LoadArchiveForm, ReactiveUploadForm
from laboratory.models import (
    OrganizationStructure,
    Laboratory,
    TemporalUploadReactive,
    Catalog,
    LaboratoryRoom,
    Furniture,
    Shelf,
)
from datetime import datetime

from laboratory.shelfobject.utils import move_shelfobject_partial_quantity_to
from laboratory.utils_upload_reatives import (
    get_reactive_by_cas_or_name,
    get_or_create_material,
    create_shelfobject,
    create_reactive_limits,
    get_units,
    validate_shelf,
)
from django.utils.translation import gettext as _

HEADER_ROW = 7
DATA_START_ROW = 8

COLUMNS = {
    "nombre_producto": 1,  # A
    "numero_cas": 2,  # B
    "formula_quimica": 3,  # C
    "estado": 4,  # D
    "cantidad": 5,  # E
    "unidades_cantidad": 6,  # F
    "capacidad_envase": 7,  # G
    "unidades_capacidad": 8,  # H
    "numero_envases": 9,  # I
    "material_contenedor": 10,  # J
    "cantidad_maxima_anual": 11,  # K
    "unidades_cantidad_maxima": 12,  # L
    "fecha_caducidad": 13,  # M
}


def read_xlsm_data(uploaded_file, shelf):
    wb = openpyxl.load_workbook(uploaded_file, keep_vba=True)
    ws = wb["Formato de Iventario"]

    records = []
    for row in ws.iter_rows(
        min_row=DATA_START_ROW, max_row=ws.max_row, values_only=True
    ):
        nombre = row[COLUMNS["nombre_producto"] - 1]

        if not nombre:
            continue

        estado = row[COLUMNS["estado"] - 1]
        unidades = row[COLUMNS["unidades_cantidad"] - 1]
        fecha = row[COLUMNS["fecha_caducidad"] - 1]
        numero_envases = row[COLUMNS["numero_envases"] - 1] or 1

        record = {
            "nombre_producto": nombre,
            "numero_cas": row[COLUMNS["numero_cas"] - 1],
            "formula_quimica": row[COLUMNS["formula_quimica"] - 1],
            "estado": None if estado == "Seleccione" else estado,
            "cantidad": row[COLUMNS["cantidad"] - 1],
            "unidades_cantidad": None if unidades == "Seleccione" else unidades,
            "capacidad_envase": row[COLUMNS["capacidad_envase"] - 1],
            "unidades_capacidad": row[COLUMNS["unidades_capacidad"] - 1],
            "material_contenedor": row[COLUMNS["material_contenedor"] - 1],
            "cantidad_maxima_anual": row[COLUMNS["cantidad_maxima_anual"] - 1],
            "unidades_cantidad_maxima": row[COLUMNS["unidades_cantidad_maxima"] - 1],
            "fecha_caducidad": (
                fecha.date().strftime("%d-%m-%Y")
                if isinstance(fecha, datetime)
                else fecha
            ),
            "shelf": shelf,
        }

        for _ in range(int(numero_envases)):
            records.append(record.copy())

    return records


@login_required
@permission_required("laboratory.add_shelfobject", raise_exception=True)
def load_archive(request, org_pk, lab_pk):
    org = get_object_or_404(OrganizationStructure, pk=org_pk)
    lab = get_object_or_404(Laboratory, pk=lab_pk)
    user_is_allowed_on_organization(request.user, org)
    organization_can_change_laboratory(lab, org)
    data = {}
    if request.method == "POST":

        know_places = request.POST.get("create-now_places")
        data = {
            "file": request.POST.get("create-file"),
        }
        if not know_places:
            lab_room, created = LaboratoryRoom.objects.get_or_create(
                laboratory=lab, name="Reactivos Cargados"
            )
            furniture, fu_created = Furniture.objects.get_or_create(
                labroom=lab_room,
                name="Mueble de Reactivos Cargados",
                type=Catalog.objects.filter(key="furniture_type").first(),
            )
            shelf, shelf_created = Shelf.objects.get_or_create(
                furniture=furniture,
                name="Estante de Reactivos Cargados",
                type=Catalog.objects.filter(key="container_type").first(),
            )
            if shelf_created:
                furniture.dataconfig = "[[[%d]]]" % shelf.pk
                furniture.save()

            data.update(
                {
                    "lab_room": lab_room.pk,
                    "furniture": furniture.pk,
                    "shelf": shelf.pk,
                }
            )

        else:
            data.update(
                {
                    "lab_room": request.POST.get("create-lab_room"),
                    "furniture": request.POST.get("create-furniture"),
                    "shelf": request.POST.get("create-shelf"),
                }
            )
        data = {k: v for k, v in data.items() if v}
        serializer = LoadArchiveSerializer(data=data)
        if not serializer.is_valid():
            form = LoadArchiveForm(request.POST, prefix="create")
            for field, errors in serializer.errors.items():
                form.add_error(field, errors)
            return render(
                request,
                "laboratory/load_archive/load_archive.html",
                {
                    "org_pk": org_pk,
                    "lab_pk": lab_pk,
                    "laboratory": lab_pk,
                    "create_form": form,
                    "errors": serializer.errors,
                },
            )
        uploaded_file = serializer.validated_data["file"]
        data = read_xlsm_data(uploaded_file, serializer.validated_data["shelf"].pk)
        shelf = serializer.validated_data["shelf"]
        if shelf.measurement_unit != None:
            for row in data:
                if (
                    shelf.measurement_unit.description
                    != row["unidades_cantidad"].capitalize()
                ):
                    messages.error(
                        request,
                        _(
                            "The measurement unit in some reactive is different from the shelf measurement unit %s"
                        )
                        % shelf.measurement_unit.description,
                    )
                    return render(
                        request,
                        "laboratory/load_archive/load_archive.html",
                        {
                            "org_pk": org_pk,
                            "lab_pk": lab_pk,
                            "laboratory": lab_pk,
                            "create_form": LoadArchiveForm(prefix="create"),
                        },
                    )

        temp = TemporalUploadReactive.objects.create(
            shelf=serializer.validated_data["shelf"], data=data
        )

        return redirect(
            "laboratory:load_archive_create_shelfobjects",
            org_pk=org_pk,
            lab_pk=lab_pk,
            key=temp.key,
        )
    return render(
        request,
        "laboratory/load_archive/load_archive.html",
        context={
            "org_pk": org_pk,
            "lab_pk": lab_pk,
            "laboratory": lab_pk,
            "create_form": LoadArchiveForm(prefix="create"),
        },
    )


@login_required
@permission_required("laboratory.add_shelfobject", raise_exception=True)
def upload_reactives(request, org_pk, lab_pk, key):
    temp = get_object_or_404(TemporalUploadReactive, key=key)
    if request.method == "POST":

        ReactiveFormSet = formset_factory(ReactiveUploadForm, extra=0)
        formset = ReactiveFormSet(request.POST)
        if formset.is_valid():

            laboratory = get_object_or_404(Laboratory, pk=lab_pk)
            organization = get_object_or_404(OrganizationStructure, pk=org_pk)

            for form in formset:

                obj = get_reactive_by_cas_or_name(
                    form.cleaned_data["numero_cas"],
                    form.cleaned_data["nombre_producto"],
                    form.cleaned_data["formula_quimica"],
                    organization,
                )
                create_reactive_limits(
                    laboratory,
                    obj,
                    form.cleaned_data["cantidad_maxima_anual"],
                    form.cleaned_data["unidades_cantidad_maxima"],
                )

                material = get_or_create_material(
                    form.cleaned_data["material_contenedor"],
                    form.cleaned_data["capacidad_envase"],
                    form.cleaned_data["unidades_capacidad"],
                    organization,
                )

                container_material = create_shelfobject(
                    {
                        "shelf": temp.shelf,
                        "object": material,
                        "physical_status": form.cleaned_data["estado"],
                        "quantity": 1,
                        "measurement_unit": Catalog.objects.get(
                            key="units", description="Unidades"
                        ),
                        "in_where_laboratory": laboratory,
                    },
                    org_pk,
                    request,
                )
                shelves = validate_shelf(
                    temp.shelf,
                    form.cleaned_data["unidades_cantidad"],
                    form.cleaned_data["cantidad"],
                )
                container = move_shelfobject_partial_quantity_to(
                    container_material,
                    org_pk,
                    lab_pk,
                    temp.shelf,
                    request,
                    quantity=1,
                )
                create_shelfobject(
                    {
                        "shelf": temp.shelf,
                        "object": obj,
                        "physical_status": form.cleaned_data["estado"],
                        "quantity": (
                            form.cleaned_data["cantidad"] if not shelves else shelves
                        ),
                        "measurement_unit": (
                            get_units(form.cleaned_data["unidades_cantidad"]).first()
                            if not shelves
                            else temp.shelf.measurement_unit
                        ),
                        "in_where_laboratory": laboratory,
                        "reactive_expiration_date": form.cleaned_data[
                            "fecha_caducidad"
                        ],
                        "container": container,
                    },
                    org_pk,
                    request,
                )
            temp.delete()
            messages.success(request, _("The archive was loaded successfully."))
            return redirect(
                "laboratory:load_archive",
                org_pk=org_pk,
                lab_pk=lab_pk,
            )

        else:
            return render(
                request,
                "laboratory/load_archive/create_shelfobjects.html",
                context={
                    "formset": formset,
                    "org_pk": org_pk,
                    "lab_pk": lab_pk,
                    "laboratory": lab_pk,
                },
            )
    ReactiveFormSet = formset_factory(ReactiveUploadForm, extra=0)
    formset = ReactiveFormSet(initial=temp.data)

    context = {
        "formset": formset,
        "org_pk": org_pk,
        "lab_pk": lab_pk,
        "laboratory": lab_pk,
    }
    return render(
        request,
        "laboratory/load_archive/create_shelfobjects.html",
        context=context,
    )
