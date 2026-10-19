from django.contrib.contenttypes.models import ContentType
from django.db.models import F

from laboratory.models import ObjectMaximumLimit, Object, Catalog
from laboratory.utils_base_unit import cached_units
from report.utils import get_conversion_units_to_kilograms
from typing import List, Dict, Optional, Tuple
import datetime
import math
import json
import argparse
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.core.files.base import ContentFile

from risk_management.models import EstablishmentLogs
from sga.models import DangerSubstance, DangerSubstanceCategory, DangerIndication
from django.utils import timezone
from django.utils.translation import gettext as _


# Columnas del inventario que consumen el reporte de regencia y la bitácora del
# establecimiento. Un inventario vacío debe traerlas igual: `drop_duplicates` y
# `clasificar_establecimiento` las piden por nombre.
COLUMNAS_INVENTARIO = ("nombre", "cas", "condicion_proceso", "h_codes", "cantidad_t")


def _cantidad_en_toneladas(limit, tons_unit, density):
    """Cantidad de un `ObjectMaximumLimit` expresada en toneladas."""
    quantity = limit.quantity or 0
    if tons_unit is not None and limit.measurement_unit_id == tons_unit.pk:
        return quantity
    try:
        return (
            get_conversion_units_to_kilograms(
                limit.measurement_unit_id, quantity, density
            )
            / 1000
        )
    except ZeroDivisionError:
        return 0


def get_inventory(filters=None):
    if filters is None:
        filters = {}
    dict_objs = []
    objs_max = ObjectMaximumLimit.objects.filter(**filters)
    objs = Object.objects.filter(
        pk__in=objs_max.values_list("object__pk", flat=True),
    ).distinct()

    process_conditions = Catalog.objects.filter(
        key="process_condition",
    )

    # La unidad la crea una migración de datos (laboratory 0177), así que puede
    # faltar: una base vaciada (flush), una instalación cargada desde fixtures o
    # un catálogo depurado a mano. Sin ella simplemente no hay registros en
    # toneladas; no es motivo para tumbar el reporte.
    tons_unit = Catalog.objects.filter(key="units", description="Toneladas").first()

    # Todo en memoria de una vez: los límites agrupados por objeto, las fichas
    # SGA precargadas y las unidades en un mapa. Antes cada sustancia costaba
    # media docena de consultas y la conversión de unidades otras tantas.
    process_names = dict(process_conditions.values_list("pk", "description"))
    limits_by_object = {}
    for limit in objs_max.order_by("pk"):
        limits_by_object.setdefault(limit.object_id, []).append(limit)
    objs = objs.prefetch_related("substancharacteristics_object__h_code")

    with cached_units():
        for obj in objs:
            data = {}
            total_shelfobjects = 0
            sga_char = obj.get_sga_characteristic()
            density = getattr(sga_char, "density", None) if sga_char else None
            h_codes = [h_code.code for h_code in sga_char.h_code.all()] if sga_char else []
            # Los límites del objeto se toman siempre de `objs_max`, es decir,
            # respetando los filtros pedidos (laboratorios de la organización,
            # periodo...). Partir de `ObjectMaximumLimit.objects` mezclaba
            # registros de cualquier laboratorio de cualquier organización.
            obj_limits = limits_by_object.get(obj.pk, [])
            # El último registro (mayor pk) por cada proceso de condición.
            last_by_process = {}
            for limit in obj_limits:
                if limit.process_condition_id in process_names:
                    last_by_process[limit.process_condition_id] = limit
            if last_by_process:
                for process_pk in sorted(last_by_process):
                    total_shelfobjects += _cantidad_en_toneladas(
                        last_by_process[process_pk], tons_unit, density
                    )
                    data = {
                        "nombre": obj.name,
                        "cas": obj.cas_code,
                        "cantidad_t": total_shelfobjects,
                        "h_codes": ";".join(h_codes),
                        "condicion_proceso": process_names[process_pk],
                    }
            else:
                # Sin proceso de condición, el último registro del objeto.
                without_process = [
                    limit for limit in obj_limits if limit.process_condition_id is None
                ]
                if not without_process:
                    # Solo tiene registros con una condición fuera del catálogo
                    # `process_condition`: no hay cantidad que sumar.
                    continue
                total_shelfobjects += round(
                    _cantidad_en_toneladas(without_process[-1], tons_unit, density), 3
                )

                data = {
                    "nombre": obj.name,
                    "cas": obj.cas_code,
                    "cantidad_t": total_shelfobjects,
                    "h_codes": ";".join(h_codes),
                    "condicion_proceso": "",
                }
            dict_objs.append(data)

    dataframe = pd.DataFrame(dict_objs, columns=list(COLUMNAS_INVENTARIO))
    if not dataframe.empty:
        dataframe = dataframe.groupby(
            ["nombre", "cas", "condicion_proceso", "h_codes"], as_index=False
        ).agg(cantidad_t=("cantidad_t", "sum"))
    return dataframe


_MAPA_TIPO_SGA = {
    "peligros físicos": "Físico",
    "peligros para la salud": "Salud",
    "peligros para el medio ambiente": "Ambiental",
}


def cargar_sga_referencia() -> Dict[str, str]:
    danger_indication = list(
        DangerIndication.objects.all()
        .annotate(Codigo_H=F("code"), Tipo=F("danger_type"))
        .values("Codigo_H", "Tipo")
        .distinct()
    )
    df_h = pd.DataFrame(danger_indication)
    df_h["Codigo_H"] = df_h["Codigo_H"].astype(str).str.upper().str.strip()
    df_h["Tipo_normalizado"] = df_h["Tipo"]
    return dict(zip(df_h["Codigo_H"], df_h["Tipo_normalizado"]))


def normalizar_tipo_sga(tipo_raw: str) -> str:
    """Convierte valores del Excel SGA a claves de categoría del Cuadro 4."""
    return _MAPA_TIPO_SGA.get(tipo_raw.strip().lower(), tipo_raw.strip())


def cargar_cas_h(path_csv: str) -> Dict[str, str]:
    df = pd.read_csv(path_csv, dtype=str).fillna("")
    for col in ["cas", "h_codes"]:
        if col not in df.columns:
            raise ValueError("El catálogo CAS→H debe tener columnas: cas, h_codes")
    df["cas"] = df["cas"].str.strip()
    df["h_codes"] = df["h_codes"].astype(str).str.strip()
    return dict(zip(df["cas"], df["h_codes"]))


COLUMNAS_CUADRO3 = (
    "cas",
    "nombre",
    "umbral_t",
    "tipo_match",
    "h_codes_match",
    "nombre_patron",
    "condiciones_especiales",
)


# Los cargadores de sga (`load_danger_substances`, `load_danger_substances_category`)
# traducen el vocabulario de los CSV del Decreto 44741 al de los modelos: el tipo de
# coincidencia «h_categoria» se guarda como «type_match», «nombre_patron» como
# «patron_name», las categorías «Físico/Salud/Ambiental» como «physical/health/
# environmental» y la condición de proceso pasa a ser una entrada del Catalog. El
# algoritmo de clasificación sigue hablando el idioma de los CSV, así que hay que
# deshacer la traducción al leer de la base. Sin esto ninguna regla por categoría H ni
# por patrón de nombre se aplicaba, las sumatorias de Físico y Ambiental quedaban
# siempre en cero y ninguna condición de proceso coincidía.
TIPO_MATCH_DESDE_MODELO = {
    "type_match": "h_categoria",
    "patron_name": "nombre_patron",
}
CATEGORIA_DESDE_MODELO = {
    "physical": "Físico",
    "health": "Salud",
    "environmental": "Ambiental",
}


def cargar_cuadro3() -> pd.DataFrame:
    # Una fila por sustancia, con sus códigos H unidos por «;» como en el CSV
    # original. Con `.values("h_codes_match")` salía una fila por código H y el
    # merge por CAS duplicaba la sustancia del inventario (y su contribución).
    filas = []
    for sustancia in DangerSubstance.objects.prefetch_related("h_codes_match"):
        tipo = (sustancia.type_match or "").strip()
        filas.append(
            {
                "cas": (sustancia.cas_code or "").strip(),
                "nombre": sustancia.name or "",
                "umbral_t": sustancia.threshold,
                "tipo_match": TIPO_MATCH_DESDE_MODELO.get(tipo, tipo) or "cas",
                "h_codes_match": ";".join(
                    sorted(h.code for h in sustancia.h_codes_match.all())
                ),
                "nombre_patron": (sustancia.patron_name or "").strip(),
                "condiciones_especiales": (sustancia.especial_condition or "").strip(),
            }
        )
    # Sin filas, pd.DataFrame([]) no tendría NINGUNA columna; una organización sin
    # sustancias peligrosas cargadas es un caso normal, no un error, así que se
    # devuelve el marco vacío con sus columnas y el reporte sale sin filas.
    df = pd.DataFrame(filas, columns=COLUMNAS_CUADRO3)
    df["umbral_t"] = pd.to_numeric(df["umbral_t"], errors="coerce").fillna(math.inf)
    return df[list(COLUMNAS_CUADRO3)]


COLUMNAS_UMBRAL_H = (
    "h_code",
    "categoria",
    "seccion",
    "condicion_proceso",
    "umbral_t",
    "notas",
)


def cargar_umbral_por_H() -> pd.DataFrame:
    filas = []
    for fila in DangerSubstanceCategory.objects.select_related("process_condition"):
        categoria = (fila.category or "").strip()
        filas.append(
            {
                "h_code": (fila.h_code_id or "").strip().upper(),
                "categoria": CATEGORIA_DESDE_MODELO.get(categoria, categoria),
                "seccion": (fila.section or "").strip().upper(),
                "condicion_proceso": (
                    fila.process_condition.description.strip()
                    if fila.process_condition_id
                    else ""
                ),
                "umbral_t": fila.threshold,
                "notas": (fila.note or "").strip(),
            }
        )
    # Mismo caso que en cargar_cuadro3(): sin filas el marco no tendría columnas.
    df = pd.DataFrame(filas, columns=COLUMNAS_UMBRAL_H)
    df["umbral_t"] = pd.to_numeric(df["umbral_t"], errors="coerce")
    return df


def evaluar_cuadro4_individual(fila, umbral_H_df: pd.DataFrame) -> Dict:
    """
    Paso 4 del Decreto 44741: verifica si UNA sustancia supera algún umbral
    individual del Cuadro 4.  Si cantidad_t >= umbral_t para cualquier match
    → riesgo mayor directo.
    """
    cantidad = fila["cantidad_t"]
    hlist = expandir_hcodes(fila.get("h_codes", ""))
    condicion_sustancia = str(fila.get("condicion_proceso", "")).strip()

    if not hlist or cantidad <= 0:
        return {"supera": False}

    df_match = pd.DataFrame({"h_code": hlist}).merge(
        umbral_H_df, on="h_code", how="inner"
    )
    if df_match.empty:
        return {"supera": False}

    # Resolver condiciones de proceso (misma lógica que contribuciones_por_sustancia)
    filas_resueltas = []
    for h_code, grupo in df_match.groupby("h_code"):
        tiene_condiciones = grupo["condicion_proceso"].str.strip().ne("").any()
        if not tiene_condiciones:
            filas_resueltas.append(grupo)
        elif condicion_sustancia:
            match_cond = grupo[
                grupo["condicion_proceso"].str.strip() == condicion_sustancia
            ]
            if match_cond.empty:
                filas_resueltas.append(grupo.nsmallest(1, "umbral_t"))
            else:
                filas_resueltas.append(match_cond)
        else:
            filas_resueltas.append(grupo.nsmallest(1, "umbral_t"))

    df_resuelto = pd.concat(filas_resueltas, ignore_index=True)

    for _, row in df_resuelto.iterrows():
        if row["umbral_t"] > 0 and cantidad >= row["umbral_t"]:
            return {
                "supera": True,
                "h_code": row["h_code"],
                "umbral": float(row["umbral_t"]),
                "categoria": row["categoria"],
            }

    return {"supera": False}


def clasificar_establecimiento(inventario, c3, umbral_H_df, map_H_to_tipo):
    inv_c3 = evaluar_cuadro3(inventario, c3)

    # Paso 2: verificar Cuadro 3
    if (inv_c3["ratio_c3"] >= 1.0).any():
        detalles = []
        for _, r in inv_c3.iterrows():
            detalles.append(
                {
                    "nombre": r["nombre"],
                    "cas": r["cas"],
                    "cantidad_t": float(r["cantidad_t"]),
                    "nominada_c3": bool(r["nominada_c3"]),
                    "umbral_c3": (
                        None if pd.isna(r["umbral_c3"]) else float(r["umbral_c3"])
                    ),
                    "ratio_c3": (
                        None
                        if pd.isna(r["ratio_c3"])
                        else round(float(r["ratio_c3"]), 3)
                    ),
                    "h_codes": r["h_codes"],
                    "contribuciones": {},
                    "detalle_contribuciones": [],
                    "regla_cruzada_salud": False,
                    "advertencias": [],
                }
            )
        return {
            "clasificacion": "riesgo mayor",
            "criterio": "Al menos una sustancia nominada (Cuadro 3) cumple o supera su umbral.",
            "sumatorias_por_categoria": {},
            "detalles": detalles,
            "trazabilidad": {},
        }

    # Paso 4: verificar umbrales individuales de Cuadro 4
    for _, fila in inv_c3.iterrows():
        resultado_c4 = evaluar_cuadro4_individual(fila, umbral_H_df)
        if resultado_c4["supera"]:
            detalles = []
            for _, r in inv_c3.iterrows():
                detalles.append(
                    {
                        "nombre": r["nombre"],
                        "cas": r["cas"],
                        "cantidad_t": float(r["cantidad_t"]),
                        "nominada_c3": bool(r["nominada_c3"]),
                        "umbral_c3": (
                            None if pd.isna(r["umbral_c3"]) else float(r["umbral_c3"])
                        ),
                        "ratio_c3": (
                            None
                            if pd.isna(r["ratio_c3"])
                            else round(float(r["ratio_c3"]), 3)
                        ),
                        "h_codes": r["h_codes"],
                        "contribuciones": {},
                        "detalle_contribuciones": [],
                        "regla_cruzada_salud": False,
                        "advertencias": [],
                        "paso4_supera": (
                            r["nombre"] == fila["nombre"] and r["cas"] == fila["cas"]
                        ),
                    }
                )
            return {
                "clasificacion": "riesgo mayor",
                "criterio": (
                    f"Sustancia '{fila['nombre']}' supera umbral individual del Cuadro 4 "
                    f"({resultado_c4['h_code']}: {fila['cantidad_t']}t >= {resultado_c4['umbral']}t)."
                ),
                "sumatorias_por_categoria": {},
                "detalles": detalles,
                "trazabilidad": {},
            }

    # Paso 5: sumatoria por categoría usando Cuadro 4
    sum_por_categoria = {"Físico": 0.0, "Salud": 0.0, "Ambiental": 0.0}
    detalles = []
    advertencias_globales = []

    for _, fila in inv_c3.iterrows():
        res_sust = contribuciones_por_sustancia(fila, umbral_H_df, map_H_to_tipo)

        for cat, val in res_sust["contribuciones"].items():
            if cat in sum_por_categoria:
                sum_por_categoria[cat] += val

        # Propagar advertencias con nombre de sustancia
        for adv in res_sust.get("advertencias", []):
            advertencias_globales.append(f"{fila['nombre']}: {adv}")

        detalles.append(
            {
                "nombre": fila["nombre"],
                "cas": fila["cas"],
                "cantidad_t": float(fila["cantidad_t"]),
                "nominada_c3": bool(fila["nominada_c3"]),
                "umbral_c3": (
                    None if pd.isna(fila["umbral_c3"]) else float(fila["umbral_c3"])
                ),
                "ratio_c3": (
                    None
                    if pd.isna(fila["ratio_c3"])
                    else round(float(fila["ratio_c3"]), 3)
                ),
                "h_codes": fila["h_codes"],
                "contribuciones": res_sust["contribuciones"],
                "detalle_contribuciones": res_sust["detalle"],
                "regla_cruzada_salud": res_sust["regla_cruzada_salud"],
                "advertencias": res_sust.get("advertencias", []),
            }
        )

    # Redondear sumatorias
    for cat in sum_por_categoria:
        sum_por_categoria[cat] = round(sum_por_categoria[cat], 4)

    criterio = "Sumatoria por categorías SGA (Cuadro 4) "
    if any(v >= 1.0 for v in sum_por_categoria.values()):
        clasif = "riesgo mayor"
        criterio += "≥ 1 en al menos una categoría."
    else:
        clasif = "riesgo menor"
        criterio += "< 1 en todas las categorías."

    return {
        "clasificacion": clasif,
        "criterio": criterio,
        "sumatorias_por_categoria": sum_por_categoria,
        "detalles": detalles,
        "advertencias": advertencias_globales,
        "trazabilidad": {},
    }


def evaluar_cuadro3(inventario: pd.DataFrame, c3: pd.DataFrame) -> pd.DataFrame:
    # Paso A: match por CAS (entradas con tipo_match="cas" y CAS no vacío)
    c3_cas = c3[(c3["tipo_match"] == "cas") & (c3["cas"].str.strip() != "")]
    merged = inventario.merge(c3_cas[["cas", "umbral_t"]], on="cas", how="left")
    merged = merged.rename(columns={"umbral_t": "umbral_c3"})

    # Paso B: match por H-code (entradas con tipo_match="h_categoria")
    c3_hcat = c3[c3["tipo_match"] == "h_categoria"]
    for _, regla in c3_hcat.iterrows():
        h_requeridos = [
            h.strip().upper() for h in regla["h_codes_match"].split(";") if h.strip()
        ]
        umbral = regla["umbral_t"]
        for idx in merged.index:
            h_sustancia = expandir_hcodes(merged.at[idx, "h_codes"])
            if any(h in h_sustancia for h in h_requeridos):
                # Aplicar si no tiene umbral o el nuevo es más bajo
                actual = merged.at[idx, "umbral_c3"]
                if pd.isna(actual) or umbral < actual:
                    merged.at[idx, "umbral_c3"] = umbral

    # Paso C: match por patrón de nombre (entradas con tipo_match="nombre_patron")
    c3_nombre = c3[c3["tipo_match"] == "nombre_patron"]
    for _, regla in c3_nombre.iterrows():
        patrones = [
            p.strip().lower() for p in regla["nombre_patron"].split(";") if p.strip()
        ]
        umbral = regla["umbral_t"]
        for idx in merged.index:
            nombre_sustancia = str(merged.at[idx, "nombre"]).lower()
            if any(patron in nombre_sustancia for patron in patrones):
                actual = merged.at[idx, "umbral_c3"]
                if pd.isna(actual) or umbral < actual:
                    merged.at[idx, "umbral_c3"] = umbral
    merged["nominada_c3"] = ~merged["umbral_c3"].isna()
    merged["ratio_c3"] = merged.apply(
        lambda r: (
            (r["cantidad_t"] / r["umbral_c3"])
            if r["nominada_c3"] and r["umbral_c3"] > 0
            else float("nan")
        ),
        axis=1,
    )
    return merged


# ---------------------------------------------------------------------------
# Expansión de H-codes
# ---------------------------------------------------------------------------


def expandir_hcodes(h_codes_str: str) -> List[str]:
    if not h_codes_str:
        return []
    return [h.strip().upper() for h in h_codes_str.split(";") if h.strip()]


# ---------------------------------------------------------------------------
# Algoritmo por sustancia (no por H-code individual)
# ---------------------------------------------------------------------------


def contribuciones_por_sustancia(
    fila, umbral_H_df: pd.DataFrame, map_H_to_tipo: Dict[str, str]
) -> Dict:
    """
    Calcula las contribuciones de UNA sustancia a cada categoría.
    Retorna dict con:
      - contribuciones: {categoría: valor}
      - detalle: lista de explicaciones
      - regla_cruzada_salud: bool
      - advertencias: lista de advertencias sobre condiciones de proceso
    """
    cantidad = fila["cantidad_t"]
    hlist = expandir_hcodes(fila.get("h_codes", ""))
    condicion_sustancia = str(fila.get("condicion_proceso", "")).strip()
    resultado = {
        "contribuciones": {},
        "detalle": [],
        "regla_cruzada_salud": False,
        "advertencias": [],
    }

    if not hlist or cantidad <= 0:
        return resultado

    # Buscar coincidencias en Cuadro 4
    df_match = pd.DataFrame({"h_code": hlist}).merge(
        umbral_H_df, on="h_code", how="inner"
    )

    if df_match.empty:
        return resultado

    # Resolver condiciones de proceso: para cada h_code, quedarse con una sola fila
    filas_resueltas = []
    advertencias = []
    for h_code, grupo in df_match.groupby("h_code"):
        # pregunta si alguno es true
        tiene_condiciones = grupo["condicion_proceso"].str.strip().ne("").any()
        if not tiene_condiciones:
            # H-code sin condiciones → usar tal cual
            filas_resueltas.append(grupo)
        elif condicion_sustancia:
            # Sustancia especifica condición → filtrar
            match_cond = grupo[
                grupo["condicion_proceso"].str.strip() == condicion_sustancia
            ]
            if match_cond.empty:
                # Condición no encontrada → usar más conservador
                filas_resueltas.append(grupo.nsmallest(1, "umbral_t"))
                advertencias.append(
                    f"{h_code}: condicion '{condicion_sustancia}' no coincide, "
                    f"usando umbral mas bajo ({grupo['umbral_t'].min()}t)"
                )
            else:
                filas_resueltas.append(match_cond)
        else:
            # Sin condición especificada → más conservador + advertencia
            filas_resueltas.append(grupo.nsmallest(1, "umbral_t"))
            advertencias.append(
                f"{h_code}: condicion no especificada, "
                f"usando umbral mas bajo ({grupo['umbral_t'].min()}t)"
            )

    df_match = pd.concat(filas_resueltas, ignore_index=True)

    # Agrupar por categoría → tomar umbral mínimo por categoría
    contribuciones_directas = {}
    detalles = []

    for cat, grupo in df_match.groupby("categoria"):
        umbral_min = grupo["umbral_t"].min()
        if umbral_min > 0:
            contrib = cantidad / umbral_min
            contribuciones_directas[cat] = contrib
            h_usados = ", ".join(grupo["h_code"].tolist())
            # Las filas sin nota llegan como "": sin filtrarlas el detalle
            # terminaba en " [, ]".
            notas_usadas = (
                ", ".join(n for n in grupo["notas"].tolist() if n)
                if "notas" in grupo.columns
                else ""
            )
            detalle_str = f"{_(cat)}: {cantidad}/{umbral_min} = {contrib:.4f} (H-codes: {h_usados})"
            if notas_usadas:
                detalle_str += f" [{notas_usadas}]"
            detalles.append(detalle_str)

    # Regla de inclusión de Salud (Paso 5, Decreto 44741):
    # Todas las sustancias incluidas en Cuadro 4 deben contribuir a la sumatoria
    # de Salud, independientemente de la sección en que aparezcan.
    if "Salud" not in contribuciones_directas:
        umbral_min_global = df_match["umbral_t"].min()
        if umbral_min_global > 0:
            contrib_salud = cantidad / umbral_min_global
            contribuciones_directas["Salud"] = contrib_salud
            resultado["regla_cruzada_salud"] = True
            h_salud_sga = [h for h in hlist if map_H_to_tipo.get(h) == "Salud"]
            detalles.append(
                f"Salud (inclusión C4): {cantidad}/{umbral_min_global} = {contrib_salud:.4f} "
                f"(H-codes SGA salud: {', '.join(h_salud_sga) if h_salud_sga else 'ninguno'})"
            )

    resultado["contribuciones"] = contribuciones_directas
    resultado["detalle"] = detalles
    resultado["advertencias"] = advertencias
    return resultado


def build_establishment_xls(res):
    """Genera un archivo Excel con la misma información del JSON de clasificación."""
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="2F5496")
    center = Alignment(horizontal="center")

    # --- Hoja 1: Resumen ---
    ws_summary = wb.active
    ws_summary.title = _("Resumen")

    ws_summary.append([_("Clasificación"), res.get("clasificacion", "")])
    ws_summary.append([_("Criterio"), res.get("criterio", "")])
    ws_summary.append([])
    ws_summary.append([_("Sumatorias por Categoría")])
    ws_summary.append([_("Categoría"), _("Valor")])
    for cat, val in res.get("sumatorias_por_categoria", {}).items():
        ws_summary.append([cat, val])

    for row in ws_summary.iter_rows(min_row=1, max_row=ws_summary.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)

    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 60

    # --- Hoja 2: Detalles por Sustancia ---
    ws_det = wb.create_sheet(title=_("Detalle Sustancias"))

    headers = [
        _("Nombre"),
        _("CAS"),
        _("Cantidad (t)"),
        _("Nominada C3"),
        _("Umbral C3 (t)"),
        _("Ratio C3"),
        _("H-codes"),
        _("Físico"),
        _("Salud"),
        _("Ambiental"),
        _("Regla cruzada Salud"),
        _("Detalle Contribuciones"),
        _("Advertencias"),
    ]
    ws_det.append(headers)
    for cell in ws_det[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for det in res.get("detalles", []):
        contribuciones = det.get("contribuciones", {})
        ws_det.append(
            [
                det.get("nombre", ""),
                det.get("cas", ""),
                det.get("cantidad_t", ""),
                det.get("nominada_c3", ""),
                det.get("umbral_c3", ""),
                det.get("ratio_c3", ""),
                det.get("h_codes", ""),
                contribuciones.get("Físico", ""),
                contribuciones.get("Salud", ""),
                contribuciones.get("Ambiental", ""),
                det.get("regla_cruzada_salud", ""),
                "\n".join(det.get("detalle_contribuciones", [])),
                "\n".join(det.get("advertencias", [])),
            ]
        )

    for col_idx in range(1, len(headers) + 1):
        ws_det.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20
    ws_det.column_dimensions[openpyxl.utils.get_column_letter(12)].width = 50
    ws_det.column_dimensions[openpyxl.utils.get_column_letter(13)].width = 40

    # --- Hoja 3: Advertencias globales ---
    advertencias = res.get("advertencias", [])
    if advertencias:
        ws_adv = wb.create_sheet(title=_("Advertencias"))
        ws_adv.append([_("Advertencias")])
        ws_adv["A1"].font = Font(bold=True)
        for adv in advertencias:
            ws_adv.append([adv])
        ws_adv.column_dimensions["A"].width = 80

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def create_estableshment_logs_data(element, day, labs):
    moment = timezone.make_aware(datetime.datetime.combine(day, datetime.time.min))
    filters = {
        "object__type": 0,
        "created_at": day,
        "object__isnull": False,
        "measurement_unit__isnull": False,
    }
    if labs:
        filters.update({"laboratory__pk__in": labs})
        inv = get_inventory(filters)
        inv = inv.drop_duplicates(subset=["nombre", "h_codes", "cas", "cantidad_t"])
        c3 = cargar_cuadro3()
        c4 = cargar_umbral_por_H()
        mapH_tipo = cargar_sga_referencia()
        ct = ContentType.objects.filter(
            app_label=element._meta.app_label,
            model=element._meta.model_name,
        ).first()
        if not inv.empty:
            res = clasificar_establecimiento(inv, c3, c4, mapH_tipo)
            sumatories = res.get("sumatorias_por_categoria", {})
            xls_bytes = build_establishment_xls(res)
            filename = f"establishment_{element.pk}_{day}.xlsx"
            log = EstablishmentLogs.objects.create(
                content_type=ct,
                object_id=element.pk,
                table_content=json.dumps(res, indent=2),
                environmental=sumatories.get("Ambiental", 0.0),
                health=sumatories.get("Salud", 0.0),
                physical=sumatories.get("Físico", 0.0),
                establishment_status=res["clasificacion"].capitalize(),
                date=moment,
            )
            log.xls_content.save(filename, ContentFile(xls_bytes), save=True)
        else:
            EstablishmentLogs.objects.create(
                content_type=ct,
                object_id=element.pk,
                environmental=0.0,
                health=0.0,
                physical=0.0,
                establishment_status="Desconocido",
                date=moment,
            )
