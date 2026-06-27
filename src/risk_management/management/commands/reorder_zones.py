from django.contrib.auth.models import User
from django.contrib.contenttypes.models import ContentType
from openpyxl import load_workbook
from django.core.management import BaseCommand

from laboratory.models import (
    Laboratory,
    OrganizationStructure,
    OrganizationStructureRelations,
    UserOrganization,
)
from risk_management.models import Buildings, RiskZone, ZoneType


class Command(BaseCommand):
    help = "Genera datos de zonas de riesgo"

    def get_or_create_zone(self, data):
        zone = RiskZone.objects.filter(
            name=data["name"], organization=self.parent_org
        ).first()
        laboratory = Laboratory.objects.filter(name=data["laboratory"]).first()
        if not zone:
            del data["laboratory"]
            zone = RiskZone.objects.create(**data)
            self.total_new_zone += 1
            self.stdout.write(self.style.SUCCESS(f"New Zone {zone.name}"))
        zone.laboratories.add(laboratory)
        zone.save()

    def create_org_lab_relation(self, org, lab):
        OrganizationStructureRelations.objects.create(
            organization=org,
            content_type=ContentType.objects.get_for_model(Laboratory),
            object_id=lab.pk,
        )

    def create_relation_org_to_org(self, org, other_parent=None):
        OrganizationStructureRelations.objects.create(
            organization=self.parent_org if other_parent is None else other_parent,
            content_type=ContentType.objects.get_for_model(OrganizationStructure),
            object_id=org.pk,
        )

    def get_or_create_lab(self, data, org):
        lab = None
        if (
            data["name"].rstrip()
            == "Laboratorio de Biotecnología de Microalgas – LABMA".rstrip()
        ):
            lab = Laboratory.objects.filter(pk=85).first()
        elif (
            data["name"].rstrip()
            == "Hospital de Especies silvestres y menores UNA (HEMS)".rstrip()
        ):
            lab = Laboratory.objects.filter(pk=121).first()
        else:
            lab = Laboratory.objects.filter(name=data["name"]).first()
        if not lab:
            lab = Laboratory.objects.create(organization=self.parent_org, **data)

            self.create_org_lab_relation(org, lab)
            self.total_new_lab += 1
            self.stdout.write(self.style.SUCCESS(f"New Lab {lab.name}"))
        return lab

    def get_or_create_organization(self, name, other_parent=None):
        org = OrganizationStructure.objects.filter(name=name).first()
        if not org:
            org = OrganizationStructure.objects.create(
                name=name,
                parent=self.parent_org if other_parent is None else other_parent,
            )
            self.create_relation_org_to_org(org, other_parent=other_parent)
            self.total_new_org += 1
            self.stdout.write(self.style.SUCCESS(f"New Organization {org.name}"))
            UserOrganization.objects.create(
                organization=org,
                user=self.default_user,
                type_in_organization=UserOrganization.ADMINISTRATOR,
            )
        return org

    def handle(self, *args, **options):
        wb = load_workbook(
            filename="./risk_management/management/commands/111 Inventarios enviados(1)-2.xlsx"
        )
        self.total_new_zone = 0
        self.total_new_org = 0
        self.total_new_lab = 0

        ws = wb.active
        self.parent_org = OrganizationStructure.objects.get(name="UNA")
        self.default_user = User.objects.get(username="wendy.umana.herrera")
        zone_type = ZoneType.objects.get(name="Otras Zonas")

        for row in ws.iter_rows(min_row=2):
            org = self.get_or_create_organization(row[3].value)
            data_lab = {
                "name": row[0].value.rstrip(),
                "faculty_dispatch": row[2].value if row[2].value else "",
            }

            lab = self.get_or_create_lab(data_lab, org)
            if row[1].value:
                org_child = self.get_or_create_organization(
                    row[1].value, other_parent=org
                )
                self.create_relation_org_to_org(org_child, other_parent=org)
                self.create_org_lab_relation(org_child, lab)

            data_zone = {
                "name": row[4].value,
                "num_workers": 1,
                "priority": 1,
                "zone_type": zone_type,
                "organization": self.parent_org,
                "laboratory": lab,
            }
            if row[4].value:
                self.get_or_create_zone(data_zone)
        self.stdout.write(self.style.SUCCESS(f"Total new zones: {self.total_new_zone}"))
        self.stdout.write(self.style.SUCCESS(f"Total new orgs: {self.total_new_org}"))
        self.stdout.write(self.style.SUCCESS(f"Total new labs: {self.total_new_lab}"))
