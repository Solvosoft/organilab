from django.contrib.auth.models import User
from django.contrib.contenttypes.models import ContentType
from openpyxl import load_workbook
from django.core.management import BaseCommand

from laboratory.models import (
    Laboratory,
    OrganizationStructure,
    OrganizationStructureRelations,
)
from risk_management.models import Buildings, RiskZone, ZoneType


class Command(BaseCommand):
    help = "Genera datos de zonas de riesgo"

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            type=str,
            help="archivo xlsx con datos de zonas de riesgo",
        )

    def get_or_create_building(self, data, lab):
        building = Buildings.objects.filter(name=data["name"], laboratories=lab).first()
        if not building:
            building = Buildings.objects.create(**data)
            building.laboratories.add(lab)
            building.save()
        return building

    def get_or_create_zone(self, data, building):
        zone = RiskZone.objects.filter(
            name=data["name"], buildings=building, organization=data["organization"]
        ).first()
        if not zone:
            zone = RiskZone.objects.create(**data)
            if building:
                zone.buildings.add(building)
                zone.save()
        return zone

    def create_org_lab_relation(self, org, lab):
        OrganizationStructureRelations.objects.create(
            organization=org,
            content_type=ContentType.objects.get_for_model(Laboratory),
            object_id=lab.pk,
        )

    def create_relation_org_to_org(self, org):
        OrganizationStructureRelations.objects.create(
            organization=self.parent_org,
            content_type=ContentType.objects.get_for_model(OrganizationStructure),
            object_id=org.pk,
        )

    def get_or_create_lab(self, data, org):
        lab = Laboratory.objects.filter(name=data["name"]).first()
        if not lab:
            lab = Laboratory.objects.create(organization=self.parent_org, **data)

            self.create_org_lab_relation(org, lab)
        return lab

    def get_or_create_organization(self, name):
        org = OrganizationStructure.objects.filter(name=name).first()
        if not org:
            org = OrganizationStructure.objects.create(
                name=name, parent=self.parent_org
            )
            self.create_relation_org_to_org(org)
        return org

    def handle(self, *args, **options):
        filename = options["output"]
        wb = load_workbook(filename=filename)
        ws = wb.active
        self.parent_org = OrganizationStructure.objects.get(name="UNA")
        self.default_user = User.objects.get(username="wendy.umana.herrera")
        zone_type = ZoneType.objects.get(name="Otras Zonas")

        for row in ws.iter_rows(min_row=2):
            org = self.get_or_create_organization(row[3].value)
            data_lab = {
                "name": row[0].value,
                "faculty_dispatch": row[2].value,
                "geolocation": f"{row[5].value},{row[6].value}",
            }
            lab = self.get_or_create_lab(data_lab, org)

            data_building = {
                "name": row[1].value,
                "organization": org,
                "geolocation": f"{row[5].value},{row[6].value}",
                "manager": self.default_user,
            }
            building = None
            if row[1].value:
                building = self.get_or_create_building(data_building, lab)

            data_zone = {
                "name": row[4].value,
                "num_workers": 1,
                "priority": 1,
                "zone_type": zone_type,
                "organization": org,
            }
            if row[4].value:
                self.get_or_create_zone(data_zone, building)
