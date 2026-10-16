from io import BytesIO

import ezodf
import xlwt
from django.core.files.base import ContentFile
from django.db.models import Min
from django.utils import timezone
from django.utils.translation import gettext as _

from laboratory.models import ObjectMaximumLimit, OrganizationStructure
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from report.utils import (
    format_date,
    get_report_name,
)
from risk_management.utils_risk import (
    cargar_cuadro3,
    cargar_sga_referencia,
    clasificar_establecimiento,
    cargar_umbral_por_H,
    get_inventory,
)


def get_regency_filters(report):
    """Filtro del inventario, el mismo para XLS, XLSX y ODS.

    Toma los laboratorios elegidos (o todos los de la organización) y el
    `period` del formulario; sin periodo, o con una fecha que no se entiende,
    se usan todos los registros hasta hoy.
    """
    organization = OrganizationStructure.objects.filter(
        pk=report.data["organization"]
    ).first()
    filters = {
        "object__type": 0,
        "object__isnull": False,
        "measurement_unit__isnull": False,
        "laboratory__pk__in": organization.get_my_laboratories,
    }
    if report.data.get("laboratory"):
        filters.update({"laboratory__pk__in": report.data["laboratory"]})
    period = get_regency_period(report)
    if period:
        filters.update({"created_at__range": list(period)})
    return filters


def get_regency_period(report):
    """(inicio, fin) del `period` del formulario, o None si viene vacío o no se entiende."""
    dates = (report.data.get("period") or "").split("-")
    if len(dates) == 2:
        start, end = format_date(dates[0].strip()), format_date(dates[1].strip())
        if start and end:
            return start, end
    return None


def get_regency_period_label(report):
    """El período que cubren los datos, para mostrarlo en el resumen.

    Sin período elegido se usan todos los registros: va desde el primero
    hasta hoy.
    """
    period = get_regency_period(report)
    if period:
        start, end = period
    else:
        start = ObjectMaximumLimit.objects.filter(
            **get_regency_filters(report)
        ).aggregate(first=Min("created_at"))["first"]
        end = timezone.localdate()
        if start is None:
            return "-"
    return "%s - %s" % (start.strftime("%d/%m/%Y"), end.strftime("%d/%m/%Y"))


def get_regency_inventory(report):
    inv = get_inventory(get_regency_filters(report))
    return inv.drop_duplicates(subset=["nombre", "h_codes", "cas", "cantidad_t"])


def classify_regency(inv):
    c3 = cargar_cuadro3()
    c4 = cargar_umbral_por_H()
    mapH_tipo = cargar_sga_referencia()
    return clasificar_establecimiento(inv, c3, c4, mapH_tipo)


REGENCY_SHEETS = ("Resumen", "Sumatoria_categoría", "Detalle")
#: Ancho máximo de columna por hoja: el criterio del resumen es una frase que
#: debe verse entera (XLS no agranda la fila); el detalle se ajusta en líneas.
REGENCY_MAX_WIDTH = {"Resumen": 90}


def get_regency_detail_header():
    return [
        _("Name"),
        _("CAS"),
        _("Quantity"),
        _("Nominated"),
        _("Threshold"),
        _("Ratio"),
        _("H-codes"),
        _("Cross rule"),
        _("Detail"),
        _("Warnings"),
        _("Health contributions"),
        _("Physical contributions"),
        _("Environmental contributions"),
    ]


def get_regency_detail_row(details):
    return [
        details["nombre"],
        details["cas"],
        round(details["cantidad_t"], 3),
        _("Yes") if details["nominada_c3"] else _("No"),
        details["umbral_c3"],
        details["ratio_c3"],
        details["h_codes"],
        _("Yes") if details["regla_cruzada_salud"] else _("No"),
        " | ".join(details["detalle_contribuciones"] or []),
        " | ".join(details["advertencias"] or []),
        round(details["contribuciones"].get("Salud", 0.0), 3),
        round(details["contribuciones"].get("Físico", 0.0), 3),
        round(details["contribuciones"].get("Ambiental", 0.0), 3),
    ]


def get_regency_sheets(report):
    """Contenido del informe regencial, igual para XLSX, XLS y ODS.

    Devuelve ``(hojas, total)``: cada hoja es ``(nombre, filas)`` y cada fila
    ``(rol, celdas)``. El rol (title, label, header, cell, risk) dice cómo se
    dibuja; cada formato solo lo traduce a su librería, así las tres salidas
    no se desalinean. ``total`` es cuántas sustancias tiene el detalle.
    """
    title = report.data.get("title") or get_report_name(report)
    inv = get_regency_inventory(report)
    try:
        res = classify_regency(inv)
        # Las clasificaciones de «riesgo mayor» por Cuadro 3/4 no calculan
        # sumatorias y devuelven el diccionario vacío.
        sumatoria = res["sumatorias_por_categoria"]
        classification = [str(res["clasificacion"]), str(res["criterio"])]
        totals = [
            sumatoria.get("Físico", 0.0),
            sumatoria.get("Salud", 0.0),
            sumatoria.get("Ambiental", 0.0),
        ]
        detail = [("header", get_regency_detail_header())]
        detail += [("cell", get_regency_detail_row(d)) for d in res["detalles"]]
        total = len(res["detalles"])
    except Exception:
        classification = ["riesgo menor", ""]
        totals = [0.0, 0.0, 0.0]
        detail = [("cell", [_("Data not available")])]
        total = 0
    summary = [
        ("title", [title]),
        ("label", [_("Period"), get_regency_period_label(report)]),
        ("blank", []),
        ("header", [_("Classification"), _("Criterion")]),
        ("risk", classification),
    ]
    totals_sheet = [
        ("header", [_("Physical"), _("Health"), _("Environmental")]),
        ("cell", totals),
    ]
    return list(zip(REGENCY_SHEETS, (summary, totals_sheet, detail))), total


#: Estilo de cada celda según el rol de su fila (y la columna, en las filas
#: de etiqueta y de clasificación).
def regency_cell_style(role, col, value):
    if role == "label":
        return "label" if col == 0 else "text"
    if role == "risk":
        if col:
            return "cell"
        return "risk_high" if "mayor" in str(value) else "risk_low"
    return role


def regency_column_widths(name, rows, min_width=10):
    """Ancho de cada columna en caracteres; el título no cuenta."""
    max_width = REGENCY_MAX_WIDTH.get(name, 50)
    widths = {}
    for role, cells in rows:
        if role == "title":
            continue
        for col, value in enumerate(cells):
            text = "" if value is None else str(value)
            widths[col] = max(widths.get(col, min_width), min(len(text) + 2, max_width))
    return widths


BRAND = "1F4E79"
BORDER_COLOR = "BFBFBF"
RISK_COLORS = {"risk_high": ("FFC7CE", "9C0006"), "risk_low": ("C6EFCE", "006100")}


def _xlsx_styles():
    side = Side(border_style="thin", color=BORDER_COLOR)
    border = Border(top=side, bottom=side, left=side, right=side)
    wrap = Alignment(vertical="top", wrap_text=True)
    font = "Arial"
    styles = {
        "title": {"font": Font(name=font, bold=True, size=14, color=BRAND)},
        "label": {"font": Font(name=font, bold=True)},
        "text": {"font": Font(name=font)},
        "header": {
            "font": Font(name=font, bold=True, color="FFFFFF"),
            "fill": PatternFill("solid", fgColor=BRAND),
            "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
            "border": border,
        },
        "cell": {"font": Font(name=font), "alignment": wrap, "border": border},
    }
    for name, (fill, color) in RISK_COLORS.items():
        styles[name] = {
            "font": Font(name=font, bold=True, color=color),
            "fill": PatternFill("solid", fgColor=fill),
            "alignment": wrap,
            "border": border,
        }
    return styles


def write_regency_xlsx(sheets):
    styles = _xlsx_styles()
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for role, cells in rows:
            ws.append(cells)
            for col, value in enumerate(cells):
                for attr, style in styles[regency_cell_style(role, col, value)].items():
                    setattr(ws.cell(ws.max_row, col + 1), attr, style)
        for col, width in regency_column_widths(name, rows).items():
            ws.column_dimensions[get_column_letter(col + 1)].width = width
        if rows and rows[0][0] == "header" and len(rows) > 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        # Al imprimir: horizontal y todas las columnas en el ancho de la página.
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    wb.save(output)
    return output


def _xls_styles(wb):
    # xlwt solo dibuja colores de su paleta: se redefinen casillas libres.
    palette = {
        "brand": BRAND,
        "border": BORDER_COLOR,
        "risk_high_fill": RISK_COLORS["risk_high"][0],
        "risk_high_font": RISK_COLORS["risk_high"][1],
        "risk_low_fill": RISK_COLORS["risk_low"][0],
        "risk_low_font": RISK_COLORS["risk_low"][1],
    }
    for index, (name, rgb) in enumerate(palette.items(), start=0x21):
        xlwt.add_palette_colour("regency_" + name, index)
        wb.set_colour_RGB(index, *bytes.fromhex(rgb))
    borders = (
        "borders: left thin, right thin, top thin, bottom thin, "
        "left_colour regency_border, right_colour regency_border, "
        "top_colour regency_border, bottom_colour regency_border;"
    )
    styles = {
        "title": xlwt.easyxf("font: bold on, height 280, colour regency_brand;"),
        "label": xlwt.easyxf("font: bold on;"),
        "text": xlwt.easyxf(""),
        "header": xlwt.easyxf(
            "font: bold on, colour white; pattern: pattern solid, fore_colour regency_brand;"
            "align: wrap on, vert centre, horiz centre; " + borders
        ),
        "cell": xlwt.easyxf("align: wrap on, vert top; " + borders),
    }
    for name in RISK_COLORS:
        styles[name] = xlwt.easyxf(
            "font: bold on, colour regency_%s_font; "
            "pattern: pattern solid, fore_colour regency_%s_fill; "
            "align: wrap on, vert top; %s" % (name, name, borders)
        )
    return styles


def write_regency_xls(sheets):
    wb = xlwt.Workbook(encoding="utf-8")
    styles = _xls_styles(wb)
    for name, rows in sheets:
        ws = wb.add_sheet(name)
        for r, (role, cells) in enumerate(rows):
            for col, value in enumerate(cells):
                ws.write(r, col, "" if value is None else value, styles[regency_cell_style(role, col, value)])
        for col, width in regency_column_widths(name, rows).items():
            ws.col(col).width = 256 * width
        if rows and rows[0][0] == "header" and len(rows) > 1:
            ws.set_panes_frozen(True)
            ws.set_horz_split_pos(1)
            ws.set_remove_splits(True)
    output = BytesIO()
    wb.save(output)
    return output


ODS_STYLE = (
    '<style:style xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
    'style:name="regency_%(name)s" style:family="table-cell" style:parent-style-name="Default">'
    '<style:table-cell-properties %(cell)s/>'
    '<style:paragraph-properties %(paragraph)s/>'
    '<style:text-properties %(text)s/>'
    '</style:style>'
)
ODS_COLUMN_STYLE = (
    '<style:style xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'style:name="regency_col_%(width)s" style:family="table-column">'
    '<style:table-column-properties style:column-width="%(cm).2fcm"/>'
    '</style:style>'
)


def _ods_styles(doc):
    border = 'fo:border="0.06pt solid #%s"' % BORDER_COLOR
    wrap = 'fo:wrap-option="wrap" style:vertical-align="top" ' + border
    bold = 'fo:font-weight="bold" style:font-weight-asian="bold" style:font-weight-complex="bold"'
    styles = {
        "title": ("", "", '%s fo:font-size="14pt" fo:color="#%s"' % (bold, BRAND)),
        "label": ("", "", bold),
        "text": ("", "", ""),
        "header": (
            'fo:background-color="#%s" fo:wrap-option="wrap" style:vertical-align="middle" %s' % (BRAND, border),
            'fo:text-align="center"',
            '%s fo:color="#FFFFFF"' % bold,
        ),
        "cell": (wrap, "", ""),
    }
    for name, (fill, color) in RISK_COLORS.items():
        styles[name] = ('fo:background-color="#%s" %s' % (fill, wrap), "", '%s fo:color="#%s"' % (bold, color))
    for name, (cell, paragraph, text) in styles.items():
        style = ODS_STYLE % {"name": name, "cell": cell, "paragraph": paragraph, "text": text}
        doc.inject_style(style)


def _ods_column_style(doc, width, injected):
    if width not in injected:
        node = ezodf.xmlns.fake_element(ODS_COLUMN_STYLE % {"width": width, "cm": width * 0.21})
        doc.content.automatic_styles.xmlnode.append(node.xmlnode)
        injected.add(width)
    return "regency_col_%s" % width


def write_regency_ods(sheets):
    doc = ezodf.newdoc(doctype="ods")
    _ods_styles(doc)
    injected = set()
    for name, rows in sheets:
        ncols = max([len(cells) for _role, cells in rows] + [1])
        sheet = ezodf.Sheet(name, size=(max(len(rows), 1), ncols))
        doc.sheets += sheet
        for r, (role, cells) in enumerate(rows):
            for col, value in enumerate(cells):
                cell = sheet[(r, col)]
                if value is not None and value != "":
                    cell.set_value(value)
                cell.style_name = "regency_" + regency_cell_style(role, col, value)
        for col, width in regency_column_widths(name, rows).items():
            sheet.column_info(col).style_name = _ods_column_style(doc, width, injected)
    return BytesIO(doc.tobytes())


REGENCY_WRITERS = {
    "xlsx": write_regency_xlsx,
    "xls": write_regency_xls,
    "ods": write_regency_ods,
}


def save_regency_file(report, content):
    report_name = get_report_name(report)
    report.file = ContentFile(
        content.getvalue(), name=f"{report_name}.{report.file_type}"
    )
    report.save()


def report_regency(report):
    """XLSX, XLS y ODS: las mismas hojas, filas y estilos."""
    sheets, total = get_regency_sheets(report)
    save_regency_file(report, REGENCY_WRITERS[report.file_type](sheets))
    return total
